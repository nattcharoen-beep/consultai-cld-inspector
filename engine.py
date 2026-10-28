#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ConsultAI: CLD & Position Code Extractor Engine
Analyzes Control Logic Diagram (CLD) drawings to identify IO CLD vs NOT IO CLD,
extracts position codes (CLD_No1, No2, No3), detects Model Card, and exports to Excel.
"""

import os
import re
import io
from typing import Dict, Any, List, Union
import pymupdf
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def extract_cld_from_pdf(pdf_source: Union[str, bytes], filename: str) -> Dict[str, Any]:
    """
    Extracts CLD metadata from a PDF file path or raw bytes.
    """
    if isinstance(pdf_source, bytes):
        doc = pymupdf.open(stream=pdf_source, filetype="pdf")
    else:
        doc = pymupdf.open(pdf_source)
        
    page = doc[0] # Drawings are typically 1 page
    full_text = page.get_text()
    
    # 1. Extract CLD File number
    # Priority A: From filename (e.g. ET-D-01_M3-M4_2010301.pdf -> 2010301, 2010508A)
    cld_file_match = re.search(r'(\d{7}[A-Z]?)', filename)
    if cld_file_match:
        cld_file = cld_file_match.group(1)
    else:
        # Priority B: From drawing text near FILE: or sheet title
        doc_no_match = re.search(r'FILE:\s*(\d{7}[A-Z]?)', full_text)
        if not doc_no_match:
            doc_no_match = re.search(r'\b(\d{7}[A-Z]?)\b', full_text)
        cld_file = doc_no_match.group(1) if doc_no_match else "UNKNOWN"

    # 2. Extract Position Code (e.g. 2-2-7, 2-7-A, 2-8-C)
    # The Position Code ALWAYS follows the 7-digit file number:
    # e.g. "2010301 2-2-7" at the bottom margin or inside "FILE: 2010301 2-2-7"
    pos_match = re.search(r'\b\d{7}[A-Z]?\s+(\d+)-(\d+)-([0-9A-Z]{1,2})\b', full_text)
    if not pos_match:
        pos_match = re.search(r'FILE:\s*\d{7}[A-Z]?\s+(\d+)-(\d+)-([0-9A-Z]{1,2})\b', full_text)

    # 3. Extract Model Card (e.g. IMCIS22, BRC300, etc.)
    # Located at center of top row in the I/O block
    card_match = re.search(r'\b(IM[A-Z0-9]+)\b', full_text)
    if not card_match:
        card_match = re.search(r'\b(BRC\d+)\b', full_text)
    model_card = card_match.group(1) if card_match else "N/A"

    if pos_match:
        no1, no2, no3 = pos_match.group(1), pos_match.group(2), pos_match.group(3)
        status = "IO CLD"
    else:
        no1, no2, no3 = "N/A", "N/A", "N/A"
        status = "NOT IO CLD"
        model_card = "N/A" # Model card is only relevant for IO CLD

    return {
        "Source File": filename,
        "CLD file": cld_file,
        "CLD_No1": no1,
        "CLD_No2": no2,
        "CLD_No3": no3,
        "Status": status,
        "Model Card": model_card
    }

def process_multiple_pdfs(file_items: List[tuple]) -> List[Dict[str, Any]]:
    """
    file_items: list of (filename, file_path_or_bytes)
    """
    results = []
    for fname, source in file_items:
        try:
            res = extract_cld_from_pdf(source, fname)
            results.append(res)
        except Exception as e:
            results.append({
                "Source File": fname,
                "CLD file": "ERROR",
                "CLD_No1": "N/A",
                "CLD_No2": "N/A",
                "CLD_No3": "N/A",
                "Status": "ERROR",
                "Model Card": "N/A",
                "Error": str(e)
            })
    return results

def generate_excel_report(data: List[Dict[str, Any]], output_path: str = None) -> Union[str, bytes]:
    """
    Generates a professionally styled Excel workbook matching the expected deliverable.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result"
    
    # Headers
    headers = ["Source File", "CLD file", "CLD_No1", "CLD_No2", "CLD_No3", "Status", "Model Card"]
    ws.append(headers)
    
    # Header Styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    
    # Rows
    io_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # soft green
    not_io_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # soft gray
    
    for row_idx, item in enumerate(data, start=2):
        row_vals = [
            item.get("Source File", ""),
            item.get("CLD file", ""),
            item.get("CLD_No1", ""),
            item.get("CLD_No2", ""),
            item.get("CLD_No3", ""),
            item.get("Status", ""),
            item.get("Model Card", "N/A")
        ]
        ws.append(row_vals)
        
        is_io = item.get("Status") == "IO CLD"
        current_fill = io_fill if is_io else not_io_fill
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            
            # Alignments
            if col_num in [1]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
            # Status highlight
            if col_num == 6:
                cell.fill = current_fill
                if is_io:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="276A3C")
                else:
                    cell.font = Font(name="Calibri", size=10, color="595959")
    
    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    ws.row_dimensions[1].height = 26
    for r in range(2, len(data) + 2):
        ws.row_dimensions[r].height = 20

    if output_path:
        wb.save(output_path)
        return output_path
    else:
        stream = io.BytesIO()
        wb.save(stream)
        return stream.getvalue()

if __name__ == "__main__":
    # Test on local extracted directory
    test_dir = r"C:\Users\nattc\OneDrive\Desktop\AI ลองไปเรื่อย\งานของเพื่อนปั๊ม\consultai"
    test_files = sorted([f for f in os.listdir(test_dir) if f.endswith(".pdf")])
    items = [(f, os.path.join(test_dir, f)) for f in test_files]
    res = process_multiple_pdfs(items)
    out_excel = os.path.join(test_dir, "Result_Test_Output.xlsx")
    generate_excel_report(res, out_excel)
    print(f"Generated test excel at {out_excel}")
